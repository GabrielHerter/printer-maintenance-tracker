import csv
from datetime import date
import io
from xml.sax.saxutils import escape as xml_escape
import zipfile

from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db.models import Avg, Count, Q, QuerySet
from django.db.models.functions import Coalesce
from django.http import HttpRequest, HttpResponse, JsonResponse, QueryDict
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views import View
from django.views.decorators.http import require_GET, require_http_methods

from cadastros.models import ModeloImpressora, NumeroSerie, Setor, TipoDefeito
from .forms import ManutencaoForm
from .models import Manutencao, SituacaoManutencao

DEFAULT_PAGE_SIZE = 15


def _is_htmx(request: HttpRequest) -> bool:
    return bool(getattr(request, 'htmx', False) or request.headers.get('HX-Request') == 'true')


def _parse_date_param(value: str | None) -> date | None:
    if not value:
        return None
    value = value.strip()
    parsed = parse_date(value)
    if parsed:
        return parsed
    # Suporte para formato DD/MM/YYYY caso enviado manualmente
    parts = value.split('/')
    if len(parts) == 3:
        try:
            return date(int(parts[2]), int(parts[1]), int(parts[0]))
        except ValueError:
            return None
    return None


def _get_base_queryset() -> QuerySet[Manutencao]:
    return Manutencao.objects.select_related(
        'modelo',
        'numero_serie',
        'setor',
        'defeito',
    ).order_by('-data_abertura', '-criado_em')


def _apply_filters(queryset: QuerySet[Manutencao], params) -> tuple[QuerySet[Manutencao], dict]:
    situacao = (params.get('situacao') or '').strip().upper()
    setor_id = (params.get('setor') or '').strip()
    modelo_id = (params.get('modelo') or '').strip()
    defeito_id = (params.get('defeito') or '').strip()

    # Suporte flexível para busca por número de série (q, serial ou numero_serie)
    busca_serial = (
        params.get('q')
        or params.get('serial')
        or params.get('numero_serie')
        or ''
    ).strip()

    data_inicio_raw = (params.get('data_inicio') or params.get('data_de') or '').strip()
    data_fim_raw = (params.get('data_fim') or params.get('data_ate') or '').strip()

    if situacao in (SituacaoManutencao.MANUTENCAO, SituacaoManutencao.RESOLVIDO):
        queryset = queryset.filter(situacao=situacao)
    else:
        situacao = ''

    if setor_id.isdigit():
        queryset = queryset.filter(setor_id=int(setor_id))
    else:
        setor_id = ''

    if modelo_id.isdigit():
        queryset = queryset.filter(modelo_id=int(modelo_id))
    else:
        modelo_id = ''

    if defeito_id.isdigit():
        queryset = queryset.filter(defeito_id=int(defeito_id))
    else:
        defeito_id = ''

    if busca_serial:
        if params.get('numero_serie') and busca_serial.isdigit():
            queryset = queryset.filter(
                Q(numero_serie_id=int(busca_serial))
                | Q(numero_serie__serial__icontains=busca_serial)
            )
        else:
            queryset = queryset.filter(numero_serie__serial__icontains=busca_serial)

    data_inicio = _parse_date_param(data_inicio_raw)
    if data_inicio:
        queryset = queryset.filter(data_abertura__gte=data_inicio)

    data_fim = _parse_date_param(data_fim_raw)
    if data_fim:
        queryset = queryset.filter(data_abertura__lte=data_fim)

    filtros_ativos = bool(
        situacao or setor_id or modelo_id or defeito_id or busca_serial or data_inicio or data_fim
    )

    filter_state = {
        'situacao': situacao,
        'setor_id': setor_id,
        'modelo_id': modelo_id,
        'defeito_id': defeito_id,
        'q': busca_serial,
        'data_inicio': data_inicio.strftime('%Y-%m-%d') if data_inicio else '',
        'data_fim': data_fim.strftime('%Y-%m-%d') if data_fim else '',
        'filtros_ativos': filtros_ativos,
    }
    return queryset, filter_state


def _get_stats(queryset: QuerySet[Manutencao] | None = None) -> dict:
    qs = queryset if queryset is not None else Manutencao.objects.all()
    agg = qs.aggregate(
        total=Count('id'),
        em_manutencao=Count('id', filter=Q(situacao=SituacaoManutencao.MANUTENCAO)),
        resolvidos=Count('id', filter=Q(situacao=SituacaoManutencao.RESOLVIDO)),
        tempo_medio=Avg('tempo_manutencao_dias'),
    )
    tempo_medio_val = agg['tempo_medio']
    mttr = round(float(tempo_medio_val), 1) if tempo_medio_val is not None else 0.0
    return {
        'total': agg['total'] or 0,
        'em_manutencao': agg['em_manutencao'] or 0,
        'resolvidos': agg['resolvidos'] or 0,
        'tempo_medio': mttr,
        'mttr': mttr,
    }


@require_GET
def dashboard_kpis_partial(request: HttpRequest) -> HttpResponse:
    """
    Fragmento HTMX que escuta 'manutencaoAtualizada from:body' e renderiza
    os 4 cards de indicadores (KPIs) calculados diretamente via Django ORM aggregate().
    """
    base_qs = Manutencao.objects.all()
    filtered_qs, _ = _apply_filters(base_qs, request.GET)
    stats = _get_stats(filtered_qs)
    context = {
        'stats': stats,
        'total_manutencoes': stats['total'],
        'em_manutencao': stats['em_manutencao'],
        'resolvidas': stats['resolvidos'],
        'mttr_dias': stats['tempo_medio'],
    }
    return render(request, 'manutencoes/partials/dashboard_kpis.html', context)


@require_GET
def dashboard_charts_api(request: HttpRequest) -> JsonResponse:
    """
    Endpoint JSON analítico para o Chart.js.
    Todas as agregações (Count, Avg) são executadas diretamente no PostgreSQL
    via Django ORM (aggregate, values, annotate, values_list), sem loops em Python.
    """
    base_qs = Manutencao.objects.all()
    qs, _ = _apply_filters(base_qs, request.GET)

    # 1. KPIs consolidados e Distribuição por Situação (Manutenção vs Resolvido)
    kpis_agg = qs.aggregate(
        total=Count('id'),
        em_manutencao=Count('id', filter=Q(situacao=SituacaoManutencao.MANUTENCAO)),
        resolvidos=Count('id', filter=Q(situacao=SituacaoManutencao.RESOLVIDO)),
        mttr_dias=Avg('tempo_manutencao_dias'),
    )
    mttr_val = kpis_agg['mttr_dias']
    mttr_rounded = round(float(mttr_val), 1) if mttr_val is not None else 0.0

    situacao_qs = (
        qs.values('situacao')
        .annotate(total=Count('id'), mttr_dias=Avg('tempo_manutencao_dias'))
        .order_by('situacao')
    )

    # 2. Top 5 Modelos de Impressora com mais manutenções (modelo__nome)
    top_modelos_qs = (
        qs.values('modelo__nome')
        .annotate(total=Count('id'), mttr_dias=Avg('tempo_manutencao_dias'))
        .order_by('-total', 'modelo__nome')[:5]
    )

    # 3. Top 5 Setores com maior volume de chamados (setor__nome)
    top_setores_qs = (
        qs.values('setor__nome')
        .annotate(total=Count('id'), mttr_dias=Avg('tempo_manutencao_dias'))
        .order_by('-total', 'setor__nome')[:5]
    )

    # 4. Top 5 Tipos de Defeito mais recorrentes (defeito__descricao)
    top_defeitos_qs = (
        qs.values('defeito__descricao')
        .annotate(total=Count('id'), mttr_dias=Avg('tempo_manutencao_dias'))
        .order_by('-total', 'defeito__descricao')[:5]
    )

    payload = {
        'kpis': {
            'total': kpis_agg['total'] or 0,
            'em_manutencao': kpis_agg['em_manutencao'] or 0,
            'resolvidos': kpis_agg['resolvidos'] or 0,
            'mttr_dias': mttr_rounded,
        },
        'situacao': {
            'labels': ['Em Manutenção', 'Resolvido'],
            'keys': [SituacaoManutencao.MANUTENCAO, SituacaoManutencao.RESOLVIDO],
            'data': [
                kpis_agg['em_manutencao'] or 0,
                kpis_agg['resolvidos'] or 0,
            ],
            'items': list(situacao_qs),
        },
        'top_modelos': {
            'labels': list(top_modelos_qs.values_list('modelo__nome', flat=True)),
            'data': list(top_modelos_qs.values_list('total', flat=True)),
            'items': list(top_modelos_qs),
        },
        'top_setores': {
            'labels': list(top_setores_qs.values_list('setor__nome', flat=True)),
            'data': list(top_setores_qs.values_list('total', flat=True)),
            'items': list(top_setores_qs),
        },
        'top_defeitos': {
            'labels': list(top_defeitos_qs.values_list('defeito__descricao', flat=True)),
            'data': list(top_defeitos_qs.values_list('total', flat=True)),
            'items': list(top_defeitos_qs),
        },
    }
    return JsonResponse(payload)


def _build_list_context(request: HttpRequest) -> dict:
    base_qs = _get_base_queryset()
    filtered_qs, filter_state = _apply_filters(base_qs, request.GET)
    stats = _get_stats(filtered_qs)

    per_page_raw = (request.GET.get('per_page') or '').strip()
    if per_page_raw.isdigit() and 1 <= int(per_page_raw) <= 100:
        per_page = int(per_page_raw)
    else:
        per_page = DEFAULT_PAGE_SIZE

    paginator = Paginator(filtered_qs, per_page)
    page_number = request.GET.get('page') or 1
    page_obj = paginator.get_page(page_number)
    elided_page_range = list(
        paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    )

    query_params = request.GET.copy()
    query_params.pop('page', None)
    # Remove chaves vazias para manter a query string limpa
    for key in list(query_params.keys()):
        if not str(query_params.get(key) or '').strip():
            query_params.pop(key, None)
    filter_querystring = query_params.urlencode()

    return {
        'manutencoes': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'elided_page_range': elided_page_range,
        'filter_querystring': filter_querystring,
        'per_page': per_page,
        'stats': stats,
        'total_manutencoes': stats['total'],
        'em_manutencao': stats['em_manutencao'],
        'resolvidas': stats['resolvidos'],
        'mttr_dias': stats['tempo_medio'],
        'filtros': filter_state,
        'situacao_choices': SituacaoManutencao.choices,
        'modelos_filtro': ModeloImpressora.objects.filter(ativo=True).order_by('nome'),
        'setores_filtro': Setor.objects.filter(ativo=True).order_by('nome'),
        'defeitos_filtro': TipoDefeito.objects.filter(ativo=True).order_by('descricao'),
        'hoje_iso': timezone.localdate().strftime('%Y-%m-%d'),
    }


@require_GET
def manutencao_list(request: HttpRequest) -> HttpResponse:
    context = _build_list_context(request)

    if _is_htmx(request):
        return render(request, 'manutencoes/partials/tabela_manutencoes.html', context)

    return render(request, 'manutencoes/manutencao_list.html', context)


class DashboardView(View):
    """
    Tela Analítica Gerencial (/dashboard/) focada 100% nos indicadores (KPIs)
    e nos 4 gráficos interativos do Chart.js.
    """

    template_name = 'manutencoes/dashboard.html'

    def get(self, request: HttpRequest, *args, **kwargs) -> HttpResponse:
        base_qs = Manutencao.objects.all()
        filtered_qs, filter_state = _apply_filters(base_qs, request.GET)
        stats = _get_stats(filtered_qs)

        context = {
            'stats': stats,
            'total_manutencoes': stats['total'],
            'em_manutencao': stats['em_manutencao'],
            'resolvidas': stats['resolvidos'],
            'mttr_dias': stats['tempo_medio'],
            'filtros': filter_state,
            'situacao_choices': SituacaoManutencao.choices,
            'modelos_filtro': ModeloImpressora.objects.filter(ativo=True).order_by('nome'),
            'setores_filtro': Setor.objects.filter(ativo=True).order_by('nome'),
            'defeitos_filtro': TipoDefeito.objects.filter(ativo=True).order_by('descricao'),
            'hoje_iso': timezone.localdate().strftime('%Y-%m-%d'),
        }

        if _is_htmx(request):
            return render(request, 'manutencoes/partials/dashboard_kpis.html', context)

        return render(request, self.template_name, context)


dashboard = DashboardView.as_view()


@require_http_methods(['GET', 'POST'])
def manutencao_create(request: HttpRequest) -> HttpResponse:
    if request.method == 'POST':
        form = ManutencaoForm(request.POST)
        if form.is_valid():
            form.save()

            if _is_htmx(request):
                response = HttpResponse('')
                response['HX-Trigger'] = 'manutencaoAtualizada'
                return response

            return redirect('manutencoes:manutencao_list')

        context = {
            'form': form,
            'is_edit': False,
            'modal_title': 'Nova Ordem de Manutenção',
            'action_url': request.path,
        }
        return render(request, 'manutencoes/partials/modal_form.html', context)

    form = ManutencaoForm()
    context = {
        'form': form,
        'is_edit': False,
        'modal_title': 'Nova Ordem de Manutenção',
        'action_url': request.path,
    }
    return render(request, 'manutencoes/partials/modal_form.html', context)


@require_http_methods(['GET', 'POST'])
def manutencao_update(request: HttpRequest, pk: int) -> HttpResponse:
    manutencao = get_object_or_404(_get_base_queryset(), pk=pk)

    if request.method == 'POST':
        form = ManutencaoForm(request.POST, instance=manutencao)
        if form.is_valid():
            form.save()

            if _is_htmx(request):
                response = HttpResponse('')
                response['HX-Trigger'] = 'manutencaoAtualizada'
                return response

            return redirect('manutencoes:manutencao_list')

        context = {
            'form': form,
            'manutencao': manutencao,
            'is_edit': True,
            'modal_title': f'Editar Manutenção #{manutencao.pk}',
            'action_url': request.path,
        }
        return render(request, 'manutencoes/partials/modal_form.html', context)

    form = ManutencaoForm(instance=manutencao)
    context = {
        'form': form,
        'manutencao': manutencao,
        'is_edit': True,
        'modal_title': f'Editar Manutenção #{manutencao.pk}',
        'action_url': request.path,
    }
    return render(request, 'manutencoes/partials/modal_form.html', context)


@require_http_methods(['POST', 'PUT', 'PATCH'])
def dar_baixa_manutencao(request: HttpRequest, pk: int) -> HttpResponse:
    manutencao = get_object_or_404(_get_base_queryset(), pk=pk)

    if request.method in ('PUT', 'PATCH') and not request.POST:
        data = QueryDict(request.body)
    else:
        data = request.POST

    raw_data_retorno = (data.get('data_retorno') or '').strip()
    hoje_iso = timezone.localdate().strftime('%Y-%m-%d')

    if not raw_data_retorno:
        context = {
            'manutencao': manutencao,
            'hoje_iso': hoje_iso,
            'baixa_erro': 'Informe a data de retorno.',
        }
        return render(request, 'manutencoes/partials/linha_manutencao.html', context)

    parsed_retorno = _parse_date_param(raw_data_retorno)
    if not parsed_retorno:
        context = {
            'manutencao': manutencao,
            'hoje_iso': hoje_iso,
            'baixa_erro': 'Data de retorno inválida.',
        }
        return render(request, 'manutencoes/partials/linha_manutencao.html', context)

    manutencao.data_retorno = parsed_retorno
    try:
        manutencao.save()
    except ValidationError as exc:
        # Restaura estado original para exibição limpa da linha com o erro
        manutencao.refresh_from_db()
        mensagens = []
        if hasattr(exc, 'message_dict'):
            for msgs in exc.message_dict.values():
                mensagens.extend(msgs)
        else:
            mensagens.extend(exc.messages)
        context = {
            'manutencao': manutencao,
            'hoje_iso': hoje_iso,
            'baixa_erro': ' '.join(mensagens) or 'Erro ao registrar data de retorno.',
        }
        return render(request, 'manutencoes/partials/linha_manutencao.html', context)

    manutencao = _get_base_queryset().get(pk=manutencao.pk)
    stats = _get_stats()

    context = {
        'manutencao': manutencao,
        'stats': stats,
        'hoje_iso': hoje_iso,
        'oob_stats': False,
    }
    response = render(request, 'manutencoes/partials/linha_manutencao.html', context)
    response['HX-Trigger'] = 'manutencaoAtualizada'
    return response


@require_http_methods(['POST', 'DELETE'])
def manutencao_delete(request: HttpRequest, pk: int) -> HttpResponse:
    manutencao = get_object_or_404(Manutencao, pk=pk)
    manutencao.delete()

    if _is_htmx(request):
        response = HttpResponse('')
        response['HX-Trigger'] = 'manutencaoAtualizada'
        return response

    return redirect('manutencoes:manutencao_list')


EXPORT_HEADERS = [
    'ID',
    'Modelo',
    'Número de Série',
    'Setor',
    'Defeito',
    'Data de Abertura',
    'Data de Retorno',
    'Situação',
    'Tempo de Manutenção (Dias)',
    'Observação',
]


def _serialize_manutencao_row(m: Manutencao) -> list:
    tempo_dias = (
        m.tempo_manutencao_dias
        if m.situacao == SituacaoManutencao.RESOLVIDO and m.tempo_manutencao_dias is not None
        else m.dias_em_aberto
    )
    return [
        m.pk,
        m.modelo.nome,
        m.numero_serie.serial,
        m.setor.nome,
        m.defeito.descricao,
        m.data_abertura.strftime('%d/%m/%Y') if m.data_abertura else '',
        m.data_retorno.strftime('%d/%m/%Y') if m.data_retorno else '',
        m.get_situacao_display(),
        tempo_dias,
        m.observacao or '',
    ]


def _col_letter(col_idx: int) -> str:
    """Converte índice 1-based em coluna Excel (1 -> A, 26 -> Z, 27 -> AA)."""
    letters = []
    while col_idx > 0:
        col_idx, rem = divmod(col_idx - 1, 26)
        letters.append(chr(65 + rem))
    return ''.join(reversed(letters))


def _build_xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    """
    Gera arquivo Excel (.xlsx) válido em memória.
    Utiliza openpyxl se instalado no ambiente; caso contrário, gera pacote OOXML nativo
    com cabeçalho estilizado e tipagem numérica/texto sem dependências extras.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = 'Manutenções'
        ws.append(headers)

        header_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical='center')

        for row in rows:
            ws.append(row)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    except ImportError:
        pass

    all_rows = [headers] + rows
    sheet_rows_xml = []
    for r_idx, row_values in enumerate(all_rows, start=1):
        cells_xml = []
        style_attr = ' s="1"' if r_idx == 1 else ''
        for c_idx, val in enumerate(row_values, start=1):
            ref = f'{_col_letter(c_idx)}{r_idx}'
            if isinstance(val, (int, float)) and not isinstance(val, bool) and r_idx > 1:
                cells_xml.append(f'<c r="{ref}" t="n"><v>{val}</v></c>')
            else:
                text = xml_escape(str(val if val is not None else ''))
                cells_xml.append(
                    f'<c r="{ref}" t="inlineStr"{style_attr}><is><t xml:space="preserve">{text}</t></is></c>'
                )
        sheet_rows_xml.append(f'<row r="{r_idx}">{"".join(cells_xml)}</row>')

    worksheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{"".join(sheet_rows_xml)}</sheetData>'
        '</worksheet>'
    )
    content_types_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '</Types>'
    )
    rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/>'
        '</Relationships>'
    )
    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Manutencoes" sheetId="1" r:id="rId1"/></sheets>'
        '</workbook>'
    )
    workbook_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
        'Target="styles.xml"/>'
        '</Relationships>'
    )
    styles_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="2">'
        '<font><sz val="11"/><name val="Calibri"/></font>'
        '<font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Calibri"/></font>'
        '</fonts>'
        '<fills count="3">'
        '<fill><patternFill patternType="none"/></fill>'
        '<fill><patternFill patternType="gray125"/></fill>'
        '<fill><patternFill patternType="solid"><fgColor rgb="FF0F172A"/></patternFill></fill>'
        '</fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="2">'
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/>'
        '</cellXfs>'
        '</styleSheet>'
    )

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('[Content_Types].xml', content_types_xml)
        zf.writestr('_rels/.rels', rels_xml)
        zf.writestr('xl/workbook.xml', workbook_xml)
        zf.writestr('xl/_rels/workbook.xml.rels', workbook_rels_xml)
        zf.writestr('xl/styles.xml', styles_xml)
        zf.writestr('xl/worksheets/sheet1.xml', worksheet_xml)
    return buffer.getvalue()


@require_GET
def exportar_manutencoes(request: HttpRequest, formato: str | None = None) -> HttpResponse:
    """
    Exporta os registros de manutenção em formato CSV ou XLSX respeitando
    exatamente os mesmos filtros aplicados na tela principal (_apply_filters).
    """
    fmt = (formato or request.GET.get('formato') or 'csv').strip().lower()
    base_qs = _get_base_queryset()
    filtered_qs, _ = _apply_filters(base_qs, request.GET)

    rows = [_serialize_manutencao_row(m) for m in filtered_qs]
    timestamp = timezone.localdate().strftime('%Y%m%d')

    if fmt == 'xlsx':
        xlsx_bytes = _build_xlsx_bytes(EXPORT_HEADERS, rows)
        response = HttpResponse(
            xlsx_bytes,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="manutencoes_{timestamp}.xlsx"'
        return response

    # Exportação CSV (UTF-8 com BOM para abertura direta no Excel em Português)
    delimiter = ';' if request.GET.get('sep') == ';' else ';'
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="manutencoes_{timestamp}.csv"'
    response.write('\ufeff')
    writer = csv.writer(response, delimiter=delimiter)
    writer.writerow(EXPORT_HEADERS)
    writer.writerows(rows)
    return response
